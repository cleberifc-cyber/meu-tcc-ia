"""Exportador editável DOCX para o subconjunto Markdown suportado."""

from __future__ import annotations

import hashlib
import re
import tomllib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tcc_kit.formatting.markdown import parse_manuscript
from tcc_kit.formatting.profile import RuleProfile
from tcc_kit.formatting.report import BuildReport
from tcc_kit.sources.store import load_store


def load_project_metadata(project_dir: Path) -> dict[str, object]:
    return tomllib.loads((project_dir / "tcc.toml").read_text(encoding="utf-8")).get("project", {})


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, display, end):
        run._r.append(element)


def _reference_identity(record):
    data = {**record.fields, **record.user_overrides}
    authors = data.get("author") or []
    names = []
    if isinstance(authors, list):
        for author in authors:
            if isinstance(author, dict):
                names.append(str(author.get("family") or author.get("name") or author.get("given") or "Autor não identificado"))
    elif isinstance(authors, str):
        names.append(authors)
    issued = data.get("issued") or []
    year = str(issued[0] if isinstance(issued, list) and issued else issued or "s.d.")[:4]
    return data, names, year


def _citation_text(key: str, references: dict[str, object]) -> str:
    record = references.get(key)
    if record is None:
        return f"[citação pendente: {key}]"
    _, names, year = _reference_identity(record)
    author = names[0] if names else "Autoria não informada"
    if len(names) > 1:
        author += " et al."
    return f"({author}, {year})"


def _bibliography_text(record) -> tuple[str, str, str]:
    data, names, year = _reference_identity(record)
    author_text = "; ".join(names) if names else "Autoria não informada"
    title = str(data.get("title", "Título não informado"))
    detail = str(data.get("container_title") or data.get("publisher") or data.get("institution") or "")
    extras = []
    for field in ("volume", "issue", "pages"):
        if data.get(field):
            extras.append(str(data[field]))
    if extras:
        detail = (detail + ", " if detail else "") + ", ".join(extras)
    if data.get("doi"):
        detail = (detail + ". " if detail else "") + f"https://doi.org/{data['doi']}"
    if detail:
        detail += "."
    return author_text, title, f"{year}. {detail}".strip()


def _frontmatter_fields(blocks) -> dict[str, str]:
    result = {}
    for block in blocks:
        if block.kind != "frontmatter":
            continue
        for line in block.text.splitlines():
            match = re.fullmatch(r"\s*([A-Za-z_][A-Za-z0-9_-]*)\s*:\s*(.*?)\s*", line)
            if match:
                value = match.group(2).strip()
                if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
                    value = value[1:-1]
                result[match.group(1)] = value
    return result


def render_blocks(blocks, metadata, references, output_path: Path, profile: RuleProfile, source_hashes: dict[str, str], force: bool = False) -> BuildReport:
    report_path = output_path.with_suffix(".report.json")
    if not force and (output_path.exists() or report_path.exists()):
        raise FileExistsError(f"Saída já existe: {output_path}. Use --force para substituir.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    settings = profile.document
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(float(settings.get("margin_top_mm", 30)))
    section.left_margin = Mm(float(settings.get("margin_left_mm", 30)))
    section.bottom_margin = Mm(float(settings.get("margin_bottom_mm", 20)))
    section.right_margin = Mm(float(settings.get("margin_right_mm", 20)))
    normal = document.styles["Normal"]
    normal.font.name = str(settings.get("font_name", "Arial"))
    normal.font.size = Pt(float(settings.get("font_size_pt", 12)))
    normal.paragraph_format.line_spacing = float(settings.get("line_spacing", 1.5))
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[style_name].font.name = normal.font.name
        document.styles[style_name].font.size = Pt(12)
    _page_number(section.footer.paragraphs[0])
    frontmatter = _frontmatter_fields(blocks)
    title = frontmatter.get("title", str(metadata.get("topic", "Título provisório do TCC")))
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.add_run(title).bold = True
    for key in ("author", "institution", "course"):
        value = frontmatter.get(key, metadata.get(key))
        if value:
            p = document.add_paragraph(str(value))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    warnings = [profile.status_message(), "Citações e referências são renderizadas em formato de prévia; valide o estilo exigido pela sua instituição."]
    reference_index = {record.key: record for record in references}
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=block.level or 1)
        elif block.kind == "paragraph":
            rendered_text = re.sub(r"\[@([A-Za-z0-9_.:-]+)\]", lambda match: _citation_text(match.group(1), reference_index), block.text)
            for match in re.finditer(r"\[@([A-Za-z0-9_.:-]+)\]", block.text):
                if match.group(1) not in reference_index:
                    warnings.append(f"Citação sem registro no manuscrito: {match.group(1)}.")
            document.add_paragraph(rendered_text)
        elif block.kind == "list":
            for item in block.items or []:
                document.add_paragraph(item, style="List Bullet")
        elif block.kind == "table":
            rows = block.rows or []
            if rows:
                table = document.add_table(rows=0, cols=max(map(len, rows)))
                for row in rows:
                    cells = table.add_row().cells
                    for index, value in enumerate(row):
                        cells[index].text = value
        elif block.kind == "image" and block.image_path:
            document.add_picture(str(block.image_path), width=Mm(120))
            if block.caption:
                caption = document.add_paragraph(block.caption)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind == "diagnostic":
            warnings.append(block.text)
    if references:
        document.add_heading("Referências", level=1)
        for record in sorted(references, key=lambda item: item.key.casefold()):
            author_text, title_value, detail = _bibliography_text(record)
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(0)
            p.add_run(f"{author_text}. ")
            p.add_run(f"{title_value}. ")
            p.add_run(detail)
            effective = {**record.fields, **record.user_overrides}
            if "title" not in effective or "author" not in effective:
                warnings.append(f"Referência {record.key}: metadados de autoria/título incompletos; confira a fonte original.")
            if record.item_type not in profile.references.get("render_types", []):
                warnings.append(f"Referência {record.key}: tipo {record.item_type!r} não coberto pelo perfil.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    report = BuildReport(str(output_path), profile.id, profile.status,
                         [f"{item.get('id', 'norma')} {item.get('edition', '')}".strip() for item in profile.standards],
                         {"python-docx": "1.x"}, list(dict.fromkeys(warnings)), source_hashes)
    report.write_json(report_path)
    return report


def render_docx(project_dir: Path, output_path: Path, profile: RuleProfile, force: bool = False) -> BuildReport:
    manuscript = project_dir / "tcc.md"
    config = project_dir / "tcc.toml"
    if not manuscript.is_file():
        raise FileNotFoundError(f"Manuscrito não encontrado: {manuscript}")
    blocks = parse_manuscript(manuscript.read_text(encoding="utf-8"), project_dir)
    return render_blocks(blocks, load_project_metadata(project_dir), load_store(project_dir), output_path,
                         profile, {path.name: _hash(path) for path in (manuscript, config, project_dir / "references.json") if path.exists()}, force)
