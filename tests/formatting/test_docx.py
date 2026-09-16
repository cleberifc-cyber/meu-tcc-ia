from docx import Document

from tcc_kit.formatting.docx import render_docx
from tcc_kit.formatting.profile import RuleProfile
from tcc_kit.project import create_project
from tcc_kit.sources.models import Provenance, ReferenceRecord
from tcc_kit.sources.store import save_record


ANSWERS = {"topic": "Acessibilidade digital", "course": "Sistemas", "institution": "IFC",
           "work_type": "monografia", "problem": "Ainda não definido",
           "objectives": "Ainda não definido", "method": "Ainda não definido"}


def test_render_docx_has_headings_and_profile_warning(tmp_path):
    project = tmp_path / "project"
    create_project(project, ANSWERS)
    (project / "tcc.md").write_text("# Introdução\n\nUma afirmação [@silva2024].\n", encoding="utf-8")
    save_record(project, ReferenceRecord(
        key="silva2024", item_type="article-journal",
        fields={"author": [{"family": "Silva"}], "title": "Inovação", "container_title": "Revista Exemplo", "issued": [2024], "doi": "10.1234/x"},
        provenance={"title": Provenance("crossref", "10.1234/x", "2026-01-01T00:00:00+00:00", "https://doi.org/10.1234/x")},
        user_overrides={},
    ))
    output = tmp_path / "tcc.docx"
    report = render_docx(project, output, RuleProfile.bundled())
    document = Document(output)
    assert any(p.text == "Introdução" for p in document.paragraphs)
    assert any("Silva, 2024" in p.text for p in document.paragraphs)
    assert any("Revista Exemplo" in p.text for p in document.paragraphs)
    assert report.profile_status == "preview"
    assert report.warnings
    assert output.with_suffix(".report.json").exists()
